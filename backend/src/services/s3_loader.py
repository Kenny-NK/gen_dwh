"""Direct S3 -> PostgreSQL loader for CSV/XLSX flow runs."""

from __future__ import annotations

import asyncio
import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Awaitable, Callable

import asyncpg
import boto3
from botocore.exceptions import BotoCoreError, ClientError
from openpyxl import load_workbook
from sqlalchemy.engine import make_url

from src.core.config import settings
from src.core.db_utils import quote_ident
from src.services.record_flattening import normalize_identifier

_MAX_IDENTIFIER_LEN = 63
_SUPPORTED_EXTENSIONS = (".csv", ".xlsx")


@dataclass
class S3LoadResult:
    success: bool
    records_processed: int = 0
    source_records_total: int | None = None
    tables_processed: int = 0
    error_message: str | None = None
    state_file: str | None = None


def _normalize_headers(raw_headers: list[str]) -> list[str]:
    result: list[str] = []
    used: set[str] = set()
    for idx, header in enumerate(raw_headers):
        base = normalize_identifier(header, f"col_{idx + 1}")
        candidate = base
        suffix = 1
        while candidate in used:
            tail = f"_{suffix}"
            candidate = f"{base[: _MAX_IDENTIFIER_LEN - len(tail)]}{tail}"
            suffix += 1
        used.add(candidate)
        result.append(candidate)
    return result


def _trim_trailing_empty(values: tuple | list) -> list:
    items = list(values)
    while items and (items[-1] is None or str(items[-1]).strip() == ""):
        items.pop()
    return items


def _stringify_value(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    text = str(value).strip()
    return text or None


def _read_csv_payload(
    payload: bytes,
    row_limit: int | None = None,
) -> tuple[list[str], list[list[str | None]]]:
    decode_errors: list[str] = []
    text_content: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            text_content = payload.decode(encoding)
            break
        except UnicodeDecodeError as exc:
            decode_errors.append(f"{encoding}: {exc}")
    if text_content is None:
        raise ValueError("Не удалось декодировать CSV-файл: " + "; ".join(decode_errors))

    rows: list[list[str | None]] = []
    reader = csv.reader(io.StringIO(text_content))
    raw_header = next(reader, None)
    if raw_header is None:
        raise ValueError("CSV-файл пустой")

    normalized_headers = _normalize_headers([str(v).strip() for v in raw_header])

    for source_row in reader:
        if not source_row:
            continue
        trimmed = _trim_trailing_empty(source_row)
        if not trimmed:
            continue
        padded = trimmed + [""] * max(0, len(normalized_headers) - len(trimmed))
        row_values = [
            _stringify_value(padded[col_idx] if col_idx < len(padded) else None)
            for col_idx in range(len(normalized_headers))
        ]
        if any(value is not None for value in row_values):
            rows.append(row_values)
        if row_limit is not None and len(rows) >= row_limit:
            break
    return normalized_headers, rows


def _read_xlsx_payload(
    payload: bytes,
    row_limit: int | None = None,
) -> tuple[list[str], list[list[str | None]]]:
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)

        header_row: list[str] | None = None
        for row in iterator:
            trimmed = _trim_trailing_empty(row or ())
            if not trimmed:
                continue
            header_row = [str(cell).strip() if cell is not None else "" for cell in trimmed]
            break

        if not header_row:
            raise ValueError("XLSX-файл пустой")

        normalized_headers = _normalize_headers(header_row)
        rows: list[list[str | None]] = []

        for row in iterator:
            values = list(row or ())
            trimmed = _trim_trailing_empty(values)
            if not trimmed:
                continue
            row_values = [
                _stringify_value(trimmed[col_idx] if col_idx < len(trimmed) else None)
                for col_idx in range(len(normalized_headers))
            ]
            if any(value is not None for value in row_values):
                rows.append(row_values)
            if row_limit is not None and len(rows) >= row_limit:
                break

        return normalized_headers, rows
    finally:
        workbook.close()


def parse_s3_object_rows(
    object_key: str,
    payload: bytes,
    row_limit: int | None = None,
) -> tuple[list[str], list[list[str | None]]]:
    lower_key = object_key.lower()
    if lower_key.endswith(".csv"):
        return _read_csv_payload(payload, row_limit=row_limit)
    if lower_key.endswith(".xlsx"):
        return _read_xlsx_payload(payload, row_limit=row_limit)
    raise ValueError(
        f"Unsupported S3 file extension for '{object_key}'. Supported: csv, xlsx."
    )


def _target_connection_kwargs() -> dict[str, object]:
    url = make_url(settings.database_business_url)
    return {
        "host": url.host or "localhost",
        "port": url.port or 5432,
        "user": url.username or "gendwh",
        "password": url.password or "",
        "database": url.database or "gendwh_business",
    }


async def _ensure_table_shape(
    conn: asyncpg.Connection,
    schema_name: str,
    table_name: str,
    columns: list[str],
    upsert_column: str | None,
) -> None:
    schema_sql = quote_ident(schema_name)
    table_sql = quote_ident(table_name)
    await conn.execute(f"CREATE SCHEMA IF NOT EXISTS {schema_sql}")

    existing_rows = await conn.fetch(
        """
        SELECT column_name
        FROM information_schema.columns
        WHERE table_schema = $1 AND table_name = $2
        ORDER BY ordinal_position
        """,
        schema_name,
        table_name,
    )
    existing_columns = {str(row["column_name"]) for row in existing_rows}

    if not existing_columns:
        columns_sql = ", ".join(f"{quote_ident(column)} TEXT" for column in columns)
        await conn.execute(f"CREATE TABLE IF NOT EXISTS {schema_sql}.{table_sql} ({columns_sql})")
        existing_columns = set(columns)
    else:
        for column in columns:
            if column in existing_columns:
                continue
            await conn.execute(
                f"ALTER TABLE {schema_sql}.{table_sql} ADD COLUMN {quote_ident(column)} TEXT"
            )

    if upsert_column:
        index_name = normalize_identifier(f"uq_{table_name}_{upsert_column}", "uq_upsert")
        await conn.execute(
            f"CREATE UNIQUE INDEX IF NOT EXISTS {quote_ident(index_name)} "
            f"ON {schema_sql}.{table_sql} ({quote_ident(upsert_column)})"
        )


async def _insert_rows(
    conn: asyncpg.Connection,
    schema_name: str,
    table_name: str,
    columns: list[str],
    rows: list[list[str | None]],
    write_mode: str,
    upsert_column: str | None,
    on_chunk_inserted: Callable[[int], Awaitable[None]] | None = None,
) -> int:
    schema_sql = quote_ident(schema_name)
    table_sql = quote_ident(table_name)

    if write_mode == "replace":
        await conn.execute(f"TRUNCATE TABLE {schema_sql}.{table_sql}")

    if not rows:
        return 0

    columns_sql = ", ".join(quote_ident(column) for column in columns)
    placeholders_sql = ", ".join(f"${index}" for index in range(1, len(columns) + 1))
    insert_prefix = (
        f"INSERT INTO {schema_sql}.{table_sql} ({columns_sql}) VALUES ({placeholders_sql})"
    )

    if write_mode == "upsert" and upsert_column:
        update_columns = [column for column in columns if column != upsert_column]
        if update_columns:
            set_sql = ", ".join(
                f"{quote_ident(column)} = EXCLUDED.{quote_ident(column)}"
                for column in update_columns
            )
            statement = (
                f"{insert_prefix} ON CONFLICT ({quote_ident(upsert_column)}) DO UPDATE SET {set_sql}"
            )
        else:
            statement = f"{insert_prefix} ON CONFLICT ({quote_ident(upsert_column)}) DO NOTHING"
    else:
        statement = insert_prefix

    batch_size = 500
    inserted_total = 0
    for offset in range(0, len(rows), batch_size):
        chunk = [tuple(row) for row in rows[offset : offset + batch_size]]
        await conn.executemany(statement, chunk)
        inserted_total += len(chunk)
        if on_chunk_inserted:
            await on_chunk_inserted(inserted_total)
    return inserted_total


async def _download_s3_object(
    client, bucket_name: str, object_key: str
) -> bytes:
    def _download_sync() -> bytes:
        response = client.get_object(Bucket=bucket_name, Key=object_key)
        body = response["Body"].read()
        return bytes(body)

    return await asyncio.to_thread(_download_sync)


async def run_s3_to_postgres(
    source_config: dict,
    target_config: dict,
    tables: list[dict] | None = None,
    progress_callback: Callable[[int, int | None], Awaitable[None]] | None = None,
) -> S3LoadResult:
    """Load selected S3 CSV/XLSX objects into target PostgreSQL schema."""
    table_specs = tables or []
    if not table_specs:
        return S3LoadResult(success=False, error_message="Для потока не выбраны файлы S3")

    bucket_name = str(source_config.get("bucket") or "").strip()
    if not bucket_name:
        return S3LoadResult(success=False, error_message="Не указан S3 bucket")

    endpoint_url = str(source_config.get("aws_endpoint_url") or "").strip() or None
    client = boto3.client(
        "s3",
        aws_access_key_id=source_config.get("aws_access_key_id"),
        aws_secret_access_key=source_config.get("aws_secret_access_key"),
        endpoint_url=endpoint_url,
    )

    write_mode = str(target_config.get("write_mode") or "append").lower()
    if write_mode not in {"append", "replace", "upsert"}:
        return S3LoadResult(success=False, error_message=f"Неподдерживаемый режим записи: {write_mode}")

    target_schema = normalize_identifier(str(target_config.get("schema") or "public"), "public")
    upsert_key_raw = str(target_config.get("upsert_key") or "").strip()
    upsert_key_norm = normalize_identifier(upsert_key_raw, "id") if upsert_key_raw else None

    conn: asyncpg.Connection | None = None
    try:
        conn = await asyncpg.connect(**_target_connection_kwargs())
        total_records = 0
        total_source_records = 0
        processed_tables = 0

        for table_spec in table_specs:
            source_key = str(table_spec.get("source_key") or "").strip()
            if not source_key:
                raise ValueError("Не указан source_key для одного из S3-файлов потока")
            if not source_key.lower().endswith(_SUPPORTED_EXTENSIONS):
                raise ValueError(
                    f"Unsupported S3 file extension for '{source_key}'. Supported: csv, xlsx."
                )

            default_target = source_key.split("/")[-1].rsplit(".", 1)[0]
            target_table = normalize_identifier(
                str(table_spec.get("table_name") or default_target),
                "s3_table",
            )

            payload = await _download_s3_object(client, bucket_name, source_key)
            columns, rows = parse_s3_object_rows(source_key, payload)
            if not columns:
                raise ValueError(f"Файл '{source_key}' не содержит колонок")
            total_source_records += len(rows)
            if progress_callback:
                await progress_callback(total_records, total_source_records)

            upsert_column = None
            if write_mode == "upsert":
                if not upsert_key_norm:
                    raise ValueError("Для режима upsert требуется upsert_key")
                if upsert_key_norm not in columns:
                    raise ValueError(
                        f"Upsert key '{upsert_key_raw}' отсутствует в файле '{source_key}'"
                    )
                upsert_column = upsert_key_norm

            async with conn.transaction():
                await _ensure_table_shape(conn, target_schema, target_table, columns, upsert_column)

                async def _on_chunk_inserted(table_inserted: int) -> None:
                    if progress_callback:
                        await progress_callback(total_records + table_inserted, total_source_records)

                inserted = await _insert_rows(
                    conn,
                    target_schema,
                    target_table,
                    columns,
                    rows,
                    write_mode,
                    upsert_column,
                    on_chunk_inserted=_on_chunk_inserted,
                )

            total_records += inserted
            processed_tables += 1
            if progress_callback:
                await progress_callback(total_records, total_source_records)

        return S3LoadResult(
            success=True,
            records_processed=total_records,
            source_records_total=total_source_records,
            tables_processed=processed_tables,
        )
    except (ClientError, BotoCoreError) as exc:
        return S3LoadResult(success=False, error_message=f"S3 ошибка: {exc}")
    except Exception as exc:
        return S3LoadResult(success=False, error_message=str(exc))
    finally:
        if conn is not None:
            await conn.close()
